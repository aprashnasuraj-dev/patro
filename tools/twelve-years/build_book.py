#!/usr/bin/env python3
import os, re, io, csv, html, uuid, zipfile, shutil, hashlib, urllib.request
from pathlib import Path
from datetime import datetime, timezone
from PIL import Image
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

TITLE = "TWELVE YEARS"
SUBTITLE = "The Himalayan Tsunami"
AUTHOR = "Suraj Dahal"
PUBLISHER = "Self-published"
PUBDATE = "2026-09-27"
OUT_EPUB = "Twelve-Years-The-Himalayan-Tsunami-Suraj-Dahal.epub"
OUT_DOCX = "Twelve-Years-The-Himalayan-Tsunami-Suraj-Dahal-Illustrated.docx"

DESCRIPTION = """On the morning of 26 August 2026, the Bhotekoshi rose. At a tea stall in Timure, Dolma Tamang was counting trucks, and rupees, and the hours until her husband might say the word Qatar. In a tunnel called Adit 2, men were finishing a shift. Downstream, a wedding musician tuned his sahanai, a teacher opened his mother's Gita, and a young mother packed a tiffin she would never collect.

Then the water came.

Twelve Years is about the people the flood took, and about the people left behind to keep waiting. They stand in queues at ward offices, fill out relief forms that have no box for their grief, and keep lighting lamps for people the law will not call dead for twelve years. It was written so that the world would not look away from the families of the Trishuli–Bhotekoshi flood. It is a story of loss, love, debt and dignity, and of the terrible patience of hope."""

ABOUT = "Suraj Dahal is a writing enthusiast from Nepal. He wrote Twelve Years in the weeks after the 2026 Trishuli–Bhotekoshi flood, while families along the river were still waiting for news, in the hope that fiction might help readers understand a grief that numbers cannot hold. This is his first novel."
DEDICATION = "For every life lost, and for everyone who suffered, waited, searched, grieved, and carried on."

PARTS = [
    (1, "I", "Pending", list(range(1,12))),
    (2, "II", "Received", list(range(12,23))),
    (3, "III", "Not Registered", list(range(23,35))),
    (4, "IV", "Presumed", list(range(35,41))),
]
CHAPTER_TITLES = {
1:"Rs 1,80,000",2:"Risk Register, Rev. 7",3:"The Floor at the Back",4:"Blouses for Teej",
5:"Best Side",6:"Drawers",7:"Arjuna's Question",8:"Twenty-Four Weeks",9:"The Lalpurja",
10:"Departure",11:"Futures",12:"8:37, Several Accounts",13:"Morning Trade",14:"Adit Three",
15:"Ek, Dui, Tin",16:"The Lowest Room",17:"Al Khor",18:"Chamber, Days One to Five",
19:"Lists",20:"The Argument",21:"Video",22:"Eleven",23:"Category C",24:"The Fixer",
25:"The Ward Office",26:"Instruments",27:"Roof, Second Night",28:"Kusha",
29:"Ghewa Without Her",30:"The Watch",31:"Jeep, Night",32:"Dashain, Tihar",
33:"The Child",34:"Mingmar, Presumed",35:"The Song for When the Bride Leaves",
36:"Chainage Four Plus Three Hundred",37:"Interest",38:"Did You Send It",39:"Return",
40:"Bhadra 2095"
}
SUBJECTS = ["Bhotekoshi","Trishuli","Flood","Nepal","Disaster","Human Stories","Climate Change","Resilience","Rasuwa","Sindhupalchok","Glacial Lake Outburst","Grief","Missing Persons","Literary Fiction"]

ALT = {
1:"Book cover for Twelve Years: The Himalayan Tsunami.",
2:"Alternate cover artwork for Twelve Years: The Himalayan Tsunami.",
3:"Wraparound cover artwork for Twelve Years.",
4:"Map showing principal locations along the Bhotekoshi–Trishuli setting of the novel.",
5:"Symbolic illustration centered on the novel's twelve-year period of waiting.",
6:"Mountain and valley scene evoking the distant white peak described as washing on a line.",
7:"Illustration introducing Part I: Pending.",
8:"Illustration introducing Part II: Received.",
9:"Illustration introducing Part III: Not Registered.",
10:"Illustration introducing Part IV: Presumed.",
11:"Illustration connected with Dolma's savings passbook and the opening chapter.",
12:"An engineer's laptop and project material, illustrating the risk-register chapter.",
13:"A kite above the valley, illustrating Twenty-Four Weeks.",
14:"A sahanai associated with the wedding music in the later chapters.",
15:"A Bhagavad Gita associated with Arjuna's Question.",
16:"An illustration of the constrained lower room and its material surroundings.",
17:"An illustration associated with the accounts and records of Chapter 12.",
18:"An identification document associated with The Argument.",
19:"An underground shaft associated with Chapter 22.",
20:"Kusha sacred grass used in rites for the dead.",
21:"A ritual image associated with Ghewa Without Her.",
22:"A matchbox associated with Chainage Four Plus Three Hundred.",
23:"Tea being poured in the novel's closing sequence.",
24:"Three official papers representing the documents, forms and legal aftermath surrounding the story."
}
CAPTION = {
1:"Cover artwork.",2:"Alternate cover artwork.",3:"Wraparound cover artwork.",
4:"The river corridor and principal places of the story.",5:"Twelve years of waiting.",
6:"The white mountain above the valley.",7:"Part I — Pending.",8:"Part II — Received.",
9:"Part III — Not Registered.",10:"Part IV — Presumed.",11:"The passbook.",
12:"Risk, recorded.",13:"Twenty-four weeks.",14:"The sahanai.",15:"Arjuna's question.",
16:"The lowest room.",17:"Several accounts.",18:"Identity.",19:"The shaft.",
20:"Kusha.",21:"Ghewa.",22:"Chainage four plus three hundred.",23:"The final pouring.",
24:"Papers, categories, and waiting."
}
FALLBACK_CH = {11:1,12:2,13:8,14:35,15:7,16:16,17:12,18:20,19:22,20:28,21:29,22:36}
ANCHOR_HINTS = {
6:["washing on a line","white thing","mountain"],
11:["passbook","one lakh eighty thousand"],
12:["laptop","risk register"],
13:["kite"],
14:["sahanai"],
15:["gita","arjuna"],
16:["lowest room","brick"],
17:["8:37","several accounts"],
18:["identity","id card","citizenship"],
19:["shaft"],
20:["kusha"],
21:["ghewa"],
22:["matchbox","chainage four plus three hundred"],
23:["pour","tea"],
24:["categories in chapter 23","court procedures","invented documents"]
}

ROOT = Path("work")
INPUT = ROOT/"input"
RAW = ROOT/"raw_images"
OPT = ROOT/"optimized"
EPUB = ROOT/"epub"
OEBPS = EPUB/"OEBPS"
XHTML = OEBPS/"xhtml"
CSS = OEBPS/"css"
IMAGES = OEBPS/"images"
OUT = Path("output")
for p in [INPUT, RAW, OPT, XHTML, CSS, IMAGES, OUT]:
    p.mkdir(parents=True, exist_ok=True)

def download(url, dest):
    req=urllib.request.Request(url, headers={"User-Agent":"Mozilla/5.0"})
    with urllib.request.urlopen(req, timeout=180) as r, open(dest,"wb") as f:
        shutil.copyfileobj(r,f)

download(os.environ["DOCX_URL"], INPUT/"manuscript.docx")
download(os.environ["IMAGES_URL"], INPUT/"images.zip")

with zipfile.ZipFile(INPUT/"images.zip") as z:
    z.extractall(RAW)

image_files=[p for p in RAW.rglob("*") if p.suffix.lower() in {".jpg",".jpeg",".png",".webp"} and "contact" not in p.name.lower() and "preview" not in p.name.lower()]
groups={}
for p in image_files:
    m=re.search(r'(?i)TY[\s_.-]*(\d{2})',p.stem)
    if m:
        groups.setdefault(int(m.group(1)),[]).append(p)
if set(groups)!=set(range(1,25)):
    raise RuntimeError(f"Expected TY_01..TY_24; found: {sorted(groups)}")

def choose(cands):
    nonorig=[p for p in cands if "_orig" not in p.stem.lower() and "-orig" not in p.stem.lower()]
    pool=nonorig or cands
    scored=[]
    for p in pool:
        try:
            with Image.open(p) as im: area=im.width*im.height
        except Exception: area=0
        scored.append((area,p))
    return max(scored,key=lambda x:x[0])[1]

chosen={n:choose(groups[n]) for n in range(1,25)}
print("IMAGE INVENTORY")
for n in range(1,25): print(f"TY_{n:02d}: {chosen[n].name}")

def optimize(src, dest, cover=False):
    with Image.open(src) as im:
        im.load()
        if im.mode not in ("RGB","L"):
            if "A" in im.mode:
                bg=Image.new("RGB",im.size,"white"); bg.paste(im,mask=im.getchannel("A")); im=bg
            else: im=im.convert("RGB")
        elif im.mode=="L": im=im.convert("RGB")
        im.thumbnail((1600,2560) if cover else (1200,100000),Image.Resampling.LANCZOS)
        target=1024*1024 if cover else 500*1024
        quality=85 if cover else 83
        while True:
            buf=io.BytesIO(); im.save(buf,"JPEG",quality=quality,optimize=True,progressive=True)
            data=buf.getvalue()
            if len(data)<=target or quality<=58: break
            quality-=5
        while len(data)>target and im.width>800:
            im=im.resize((int(im.width*.9),int(im.height*.9)),Image.Resampling.LANCZOS)
            buf=io.BytesIO(); im.save(buf,"JPEG",quality=max(58,quality),optimize=True,progressive=True); data=buf.getvalue()
        dest.write_bytes(data)
        return im.size, len(data)

img_meta={}
for n,src in chosen.items():
    dest=OPT/f"ty_{n:02d}.jpg"
    dims,size=optimize(src,dest,cover=(n==1))
    shutil.copy2(dest,IMAGES/dest.name)
    img_meta[n]={"source":src.name,"file":dest.name,"dims":dims,"size":size}

doc=Document(INPUT/"manuscript.docx")
paras=list(doc.paragraphs)
def ptext(p): return re.sub(r'\s+',' ',p.text.replace("\u00a0"," ")).strip()
def next_nonempty(i):
    j=i+1
    while j<len(paras) and not ptext(paras[j]): j+=1
    return j

chapter_heading_idx={}
for n in range(1,41):
    hits=[i for i,p in enumerate(paras) if ptext(p).upper()==f"CHAPTER {n}"]
    if not hits: raise RuntimeError(f"Missing CHAPTER {n}")
    chapter_heading_idx[n]=hits[-1]
part_idx={}
for partno,roman,name,chs in PARTS:
    hits=[]
    for i,p in enumerate(paras):
        if ptext(p)==roman:
            j=next_nonempty(i)
            if j<len(paras) and ptext(paras[j]).lower()==name.lower(): hits.append(i)
    if not hits: raise RuntimeError(f"Missing Part {roman}: {name}")
    part_idx[partno]=hits[-1]
gloss_hits=[i for i,p in enumerate(paras) if ptext(p).lower()=="glossary"]
note_hits=[i for i,p in enumerate(paras) if ptext(p).lower() in {"author's note","author’s note"}]
if not gloss_hits or not note_hits: raise RuntimeError("Missing Glossary or Author's Note")
gloss_idx=max(gloss_hits); note_idx=max(note_hits)
print(f"STRUCTURE: 4 parts, {len(chapter_heading_idx)} chapters, Glossary at {gloss_idx}, Author Note at {note_idx}")

def effective_fmt(p, attr):
    v=getattr(p.paragraph_format,attr)
    if v is not None:return v
    try:return getattr(p.style.paragraph_format,attr)
    except Exception:return None
def effective_alignment(p):
    if p.alignment is not None:return p.alignment
    try:return p.style.paragraph_format.alignment
    except Exception:return None

def smart_char(s,i):
    c=s[i]; prev=s[i-1] if i else ""; nxt=s[i+1] if i+1<len(s) else ""
    if c=='"':
        return "“" if (not prev or prev.isspace() or prev in "([{\u2014\u2013") else "”"
    if c=="'":
        if prev.isalnum() and nxt.isalnum(): return "’"
        tail=s[i:i+7].lower()
        if any(tail.startswith(x) for x in ["'em","'cause","'til","'90","'80","'70","'60","'50"]): return "’"
        return "‘" if (not prev or prev.isspace() or prev in "([{\u2014\u2013") else "’"
    return c

def transformed_run_texts(p):
    runs=list(p.runs); full="".join(r.text for r in runs); k=0; out=[]
    for r in runs:
        rr=[]
        for _ in r.text:
            rr.append(smart_char(full,k)); k+=1
        out.append(re.sub(r' {2,}',' ',"".join(rr).replace("...","…")))
    return runs,out
def inline_html(p):
    runs,texts=transformed_run_texts(p); pieces=[]
    for r,t in zip(runs,texts):
        if not t: continue
        e=html.escape(t,quote=False)
        if r.italic and r.bold:e=f"<strong><em>{e}</em></strong>"
        elif r.italic:e=f"<em>{e}</em>"
        elif r.bold:e=f"<strong>{e}</strong>"
        pieces.append(e)
    return "".join(pieces) if pieces else html.escape(ptext(p),quote=False)
def plain_smart(p):
    s=ptext(p)
    return "".join(smart_char(s,i) for i in range(len(s))).replace("...","…")
def is_doc_para(p):
    left=effective_fmt(p,"left_indent"); align=effective_alignment(p)
    return bool(left and left.pt>0.5 and align==WD_ALIGN_PARAGRAPH.LEFT)
def is_noindent(p):
    fi=effective_fmt(p,"first_line_indent")
    return fi is None or abs(fi.pt)<0.5

def doc_block_html(block):
    lines=[plain_smart(p) for p in block if ptext(p)]
    flags=[(" · " in x and len(x.split(" · "))>=2) for x in lines]
    out=['<aside class="document" epub:type="z3998:document">']; i=0
    while i<len(lines):
        if flags[i]:
            rows=[]
            while i<len(lines) and flags[i]:
                rows.append([x.strip() for x in lines[i].split(" · ")]); i+=1
            if len(rows)>=2 and len({len(r) for r in rows})==1:
                header=rows[0]; looks=sum(bool(re.search(r'[A-Za-z]',c)) and not re.search(r'\d',c) for c in header)>=max(1,len(header)//2)
                out.append("<table>")
                if looks:
                    out.append("<thead><tr>"+"".join(f"<th>{html.escape(c)}</th>" for c in header)+"</tr></thead><tbody>"); data=rows[1:]
                else: out.append("<tbody>"); data=rows
                for row in data: out.append("<tr>"+"".join(f"<td>{html.escape(c)}</td>" for c in row)+"</tr>")
                out.append("</tbody></table>")
            else:
                for row in rows: out.append(f'<p>{html.escape(" · ".join(row))}</p>')
        else:
            out.append(f"<p>{html.escape(lines[i])}</p>"); i+=1
    out.append("</aside>"); return "\n".join(out)

def xhtml_shell(title, body, epub_type=None, root_rel="../"):
    et=f' epub:type="{epub_type}"' if epub_type else ""
    return f'''<?xml version="1.0" encoding="utf-8"?>
<html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops" epub:prefix="z3998: http://www.daisy.org/z3998/2012/vocab/structure/#" lang="en" xml:lang="en">
<head><meta charset="utf-8"/><title>{html.escape(title)}</title><link rel="stylesheet" type="text/css" href="{root_rel}css/styles.css"/></head>
<body{et}>{body}</body></html>'''
def fig(n):
    return f'''<figure class="illustration"><img src="../images/{img_meta[n]["file"]}" alt="{html.escape(ALT[n],quote=True)}"/><figcaption>{html.escape(CAPTION[n])}</figcaption></figure>'''

chapter_paras={}
starts=sorted((idx,n) for n,idx in chapter_heading_idx.items())
for k,(idx,n) in enumerate(starts):
    end=starts[k+1][0] if k+1<len(starts) else gloss_idx
    for pi in part_idx.values():
        if idx<pi<end:end=pi
    j=next_nonempty(idx)
    if j<end and ptext(paras[j]).lower()==CHAPTER_TITLES[n].lower(): j=next_nonempty(j)
    chapter_paras[n]=paras[j:end]

chapter_images={n:[] for n in range(1,41)}
asset_ch={}
for a in range(11,23):
    m=re.search(r'ch(?:apter)?[\s_.-]*0?(\d{1,2})',chosen[a].stem.lower())
    ch=int(m.group(1)) if m else FALLBACK_CH[a]
    if not 1<=ch<=40: ch=FALLBACK_CH[a]
    asset_ch[a]=ch; chapter_images[ch].append(a)
chapter_images[1].append(6); asset_ch[6]=1
chapter_images[40].append(23); asset_ch[23]=40

def choose_anchor(ch,a,plist):
    for i,p in enumerate(plist):
        t=ptext(p).lower()
        if any(h.lower() in t for h in ANCHOR_HINTS.get(a,[])): return i
    stem=re.sub(r'(?i)^ty[\s_.-]*\d{2}[\s_.-]*','',chosen[a].stem)
    toks=[x for x in re.split(r'[-_\s]+',stem.lower()) if len(x)>3 and not x.startswith("ch") and x not in {"orig","cover","final"}]
    for i,p in enumerate(plist):
        if any(tok in ptext(p).lower() for tok in toks):return i
    return 0
chapter_anchor={a:choose_anchor(ch,a,chapter_paras[ch]) for ch,assets in chapter_images.items() for a in assets}

placements=[]
def render_paragraphs(plist,ch=None):
    out=[]; i=0; just_broke=True; image_after={}
    if ch:
        for a in chapter_images[ch]:image_after.setdefault(chapter_anchor[a],[]).append(a)
    while i<len(plist):
        p=plist[i]; t=ptext(p)
        if not t:i+=1;continue
        if t=="*":
            out.append('<hr class="scene-break"/>');just_broke=True
            for a in image_after.get(i,[]):out.append(fig(a))
            i+=1;continue
        if is_doc_para(p):
            block=[p];j=i+1
            while j<len(plist) and (is_doc_para(plist[j]) or not ptext(plist[j])):block.append(plist[j]);j+=1
            out.append(doc_block_html(block));just_broke=True
            for idx in range(i,j):
                for a in image_after.get(idx,[]):out.append(fig(a))
            i=j;continue
        cls=[]
        if just_broke or is_noindent(p):cls.append("noindent")
        if effective_alignment(p)==WD_ALIGN_PARAGRAPH.CENTER:cls.append("center")
        c=f' class="{" ".join(cls)}"' if cls else ""
        out.append(f"<p{c}>{inline_html(p)}</p>");just_broke=False
        for a in image_after.get(i,[]):out.append(fig(a));just_broke=True
        i+=1
    return "\n".join(out)

(CSS/"styles.css").write_text(r'''body{font-family:Georgia,"Times New Roman",serif;line-height:1.5;text-align:justify;hyphens:auto;margin:5%;widows:2;orphans:2}
p{margin:0;text-indent:1.3em}p.noindent,h1+p,h2+p,hr+p,figure+p{text-indent:0}.center{text-align:center;text-indent:0}
.visually-hidden{position:absolute;width:1px;height:1px;overflow:hidden;clip:rect(0,0,0,0)}
h1,h2{font-weight:normal;text-align:center;page-break-after:avoid;break-after:avoid}.chapter{page-break-before:always;break-before:page}
.chapter-heading{margin-top:3.5em;margin-bottom:2.2em}.chapter-label{display:block;font-size:.72em;font-variant:small-caps;letter-spacing:.08em;margin-bottom:.8em}
.chapter-title{display:block;font-size:1.35em}.part-page{page-break-before:always;break-before:page;text-align:center;margin-top:35%}
.part-page h1{font-size:2em}.part-title{display:block;font-style:italic;font-size:.75em;margin-top:.8em}
.titlepage,.epigraph,.copyright,.dedication,.frontispiece,.about{text-align:center}.titlepage{margin-top:22%}.titlepage h1{font-size:2.3em;letter-spacing:.05em}
.subtitle{font-size:1.15em}.small{font-size:.85em}.epigraph{margin:22% 10% 0}.epigraph blockquote{font-style:italic}
hr.scene-break{border:0;margin:1.6em 0;text-align:center;height:1.3em}hr.scene-break:after{content:"⁂";font-size:1em}
aside.document{font-family:"Courier New",Courier,sans-serif;font-size:.88em;border-left:.18em solid #999;background:#f3f3f3;padding:.8em 1em;margin:1.2em 0}
aside.document p{text-indent:0;margin:.2em 0}table{border-collapse:collapse;width:100%;font-size:.9em;margin:.6em 0}
th,td{border:.05em solid #aaa;padding:.25em;vertical-align:top;text-align:left}th{font-weight:bold}
figure.illustration{text-align:center;margin:1.5em auto;page-break-inside:avoid;break-inside:avoid}figure.illustration img{max-width:100%;height:auto}
figcaption{font-size:.8em;color:#666;text-align:center;margin-top:.5em;text-indent:0}.cover-page{margin:0;padding:0;text-align:center}.cover-page img{max-width:100%;max-height:100%;height:auto}
nav ol{list-style-type:none;padding-left:1.2em}nav>ol{padding-left:0}nav li{margin:.35em 0}dl{margin:1em 0}dt{font-weight:bold;margin-top:.8em}dd{margin-left:1.2em}''',encoding="utf-8")

(XHTML/"cover.xhtml").write_text(xhtml_shell("Cover",f'<div class="cover-page"><img src="../images/{img_meta[1]["file"]}" alt="{html.escape(ALT[1],quote=True)}"/></div>',"cover"),encoding="utf-8")
(XHTML/"titlepage.xhtml").write_text(xhtml_shell("Title Page",f'<section class="titlepage" epub:type="titlepage"><h1>{TITLE}</h1><p class="center subtitle">{SUBTITLE}</p><p class="center">A Novel</p><p class="center">{AUTHOR}</p></section>',"titlepage"),encoding="utf-8")
(XHTML/"frontispiece.xhtml").write_text(xhtml_shell("Frontispiece",f'<section class="frontispiece"><h1 class="visually-hidden">Frontispiece</h1>{fig(2)}</section>',"frontmatter"),encoding="utf-8")
rights="Copyright © 2026 Suraj Dahal. All rights reserved."
rights2="No part of this publication may be reproduced, stored in a retrieval system, or transmitted in any form or by any means without prior written permission from the author, except for brief quotations used in reviews."
disc="This is a work of fiction inspired by the 2026 Trishuli–Bhotekoshi flood. Names, characters and incidents are products of the author's imagination; any resemblance to actual persons is coincidental."
cb=f'<section class="copyright"><h1>Copyright</h1>'+''.join(f'<p class="center small">{html.escape(t)}</p>' for t in [rights,rights2,disc,"First edition, 2026","Self-published"])+'</section>'
(XHTML/"copyright.xhtml").write_text(xhtml_shell("Copyright",cb,"copyright-page"),encoding="utf-8")
(XHTML/"dedication.xhtml").write_text(xhtml_shell("Dedication",f'<section class="dedication"><h1>Dedication</h1><p class="center"><em>{html.escape(DEDICATION)}</em></p>{fig(5)}</section>',"dedication"),encoding="utf-8")
epi='A person who has been absent and unheard of for twelve years<br/>may be presumed dead.<br/><br/>— after the Muluki Civil Code of Nepal (paraphrase)'
(XHTML/"epigraph.xhtml").write_text(xhtml_shell("Epigraph",f'<section class="epigraph"><h1 class="visually-hidden">Epigraph</h1><blockquote><p class="center">{epi}</p></blockquote></section>',"epigraph"),encoding="utf-8")
(XHTML/"map.xhtml").write_text(xhtml_shell("Map",f'<section><h1>Map</h1>{fig(4)}</section>',"frontmatter"),encoding="utf-8")

for partno,roman,name,chs in PARTS:
    (XHTML/f"part{partno:02d}.xhtml").write_text(xhtml_shell(f"Part {roman}: {name}",f'<section class="part-page"><h1>Part {roman}<span class="part-title">{html.escape(name)}</span></h1>{fig(6+partno)}</section>',"part"),encoding="utf-8")
    placements.append({"asset":6+partno,"source":chosen[6+partno].name,"placement":f"Part {roman}: {name} opener","after":"part heading"})
    for ch in chs:
        body=f'<section class="chapter" epub:type="chapter"><h1 class="visually-hidden">Part {roman}: {html.escape(name)}</h1><h2 class="chapter-heading"><span class="chapter-label">Chapter {ch}</span><span class="chapter-title">{html.escape(CHAPTER_TITLES[ch])}</span></h2>{render_paragraphs(chapter_paras[ch],ch)}</section>'
        (XHTML/f"ch{ch:02d}.xhtml").write_text(xhtml_shell(f"Chapter {ch}: {CHAPTER_TITLES[ch]}",body,"chapter"),encoding="utf-8")
        for a in chapter_images[ch]:
            idx=chapter_anchor[a]
            anchor=ptext(chapter_paras[ch][idx])[:140] if chapter_paras[ch] and idx<len(chapter_paras[ch]) else "chapter opening"
            placements.append({"asset":a,"source":chosen[a].name,"placement":f"Chapter {ch}: {CHAPTER_TITLES[ch]}","after":anchor or "chapter opening"})

placements += [
 {"asset":1,"source":chosen[1].name,"placement":"Cover","after":"official cover page"},
 {"asset":2,"source":chosen[2].name,"placement":"Frontispiece","after":"title page"},
 {"asset":4,"source":chosen[4].name,"placement":"Map","after":"Table of Contents / before Part I"},
 {"asset":5,"source":chosen[5].name,"placement":"Dedication","after":"dedication text"},
]

gloss_paras=[p for p in paras[gloss_idx+1:note_idx] if ptext(p)]
gloss_intro="";entries=[]
for p in gloss_paras:
    t=plain_smart(p)
    if not entries and "Words are Nepali unless" in t:gloss_intro=t;continue
    if " — " in t:
        term,definition=t.split(" — ",1);entries.append((term.strip(),definition.strip()))
    elif entries:
        term,definition=entries[-1];entries[-1]=(term,(definition+" "+t).strip())
    elif not gloss_intro:gloss_intro=t
gbody='<section epub:type="glossary"><h1>Glossary</h1>'
if gloss_intro:gbody+=f'<p class="noindent">{html.escape(gloss_intro)}</p>'
gbody+="<dl>"+ "".join(f"<dt>{html.escape(t)}</dt><dd>{html.escape(d)}</dd>" for t,d in entries)+"</dl></section>"
(XHTML/"glossary.xhtml").write_text(xhtml_shell("Glossary",gbody,"glossary"),encoding="utf-8")

note_paras=[p for p in paras[note_idx+1:] if ptext(p)]
(XHTML/"authors-note.xhtml").write_text(xhtml_shell("Author's Note",f'<section epub:type="afterword"><h1>Author’s Note</h1>{render_paragraphs(note_paras)}{fig(24)}</section>',"afterword"),encoding="utf-8")
placements.append({"asset":24,"source":chosen[24].name,"placement":"Author's Note","after":"author note text"})
(XHTML/"about.xhtml").write_text(xhtml_shell("About the Author",f'<section class="about" epub:type="contributors"><h1>About the Author</h1><p class="noindent">{html.escape(ABOUT)}</p>{fig(3)}</section>',"contributors"),encoding="utf-8")
placements.append({"asset":3,"source":chosen[3].name,"placement":"About the Author","after":"author biography"})

front_links=[("xhtml/titlepage.xhtml","Title Page"),("xhtml/copyright.xhtml","Copyright"),("xhtml/dedication.xhtml","Dedication"),("xhtml/epigraph.xhtml","Epigraph"),("xhtml/map.xhtml","Map")]
nav_parts=[]
for partno,roman,name,chs in PARTS:
    sub="".join(f'<li><a href="xhtml/ch{ch:02d}.xhtml">Chapter {ch}. {html.escape(CHAPTER_TITLES[ch])}</a></li>' for ch in chs)
    nav_parts.append(f'<li><a href="xhtml/part{partno:02d}.xhtml">Part {roman}: {html.escape(name)}</a><ol>{sub}</ol></li>')
nav_body=f'''<nav epub:type="toc" id="toc"><h1>Table of Contents</h1><ol>{''.join(f'<li><a href="{href}">{html.escape(label)}</a></li>' for href,label in front_links)}{''.join(nav_parts)}<li><a href="xhtml/glossary.xhtml">Glossary</a></li><li><a href="xhtml/authors-note.xhtml">Author’s Note</a></li><li><a href="xhtml/about.xhtml">About the Author</a></li></ol></nav>
<nav epub:type="landmarks" hidden="hidden"><h2>Landmarks</h2><ol><li><a epub:type="cover" href="xhtml/cover.xhtml">Cover</a></li><li><a epub:type="toc" href="nav.xhtml#toc">Table of Contents</a></li><li><a epub:type="bodymatter" href="xhtml/ch01.xhtml">Start of Story</a></li></ol></nav>'''
(OEBPS/"nav.xhtml").write_text(xhtml_shell("Table of Contents",nav_body,"toc",root_rel=""),encoding="utf-8")

play=1
def navpoint(label,src,children=""):
    global play
    n=play;play+=1
    return f'<navPoint id="navPoint-{n}" playOrder="{n}"><navLabel><text>{html.escape(label)}</text></navLabel><content src="{src}"/>{children}</navPoint>'
book_uuid=str(uuid.uuid4());nodes=[]
for href,label in front_links:nodes.append(navpoint(label,href))
for partno,roman,name,chs in PARTS:
    child="".join(navpoint(f"Chapter {ch}. {CHAPTER_TITLES[ch]}",f"xhtml/ch{ch:02d}.xhtml") for ch in chs)
    nodes.append(navpoint(f"Part {roman}: {name}",f"xhtml/part{partno:02d}.xhtml",child))
nodes += [navpoint("Glossary","xhtml/glossary.xhtml"),navpoint("Author's Note","xhtml/authors-note.xhtml"),navpoint("About the Author","xhtml/about.xhtml")]
(OEBPS/"toc.ncx").write_text(f'''<?xml version="1.0" encoding="UTF-8"?><ncx xmlns="http://www.daisy.org/z3986/2005/ncx/" version="2005-1"><head><meta name="dtb:uid" content="urn:uuid:{book_uuid}"/><meta name="dtb:depth" content="2"/><meta name="dtb:totalPageCount" content="0"/><meta name="dtb:maxPageNumber" content="0"/></head><docTitle><text>{TITLE}</text></docTitle><docAuthor><text>{AUTHOR}</text></docAuthor><navMap>{''.join(nodes)}</navMap></ncx>''',encoding="utf-8")

modified=datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00","Z")
manifest=[('nav','nav.xhtml','application/xhtml+xml','nav'),('ncx','toc.ncx','application/x-dtbncx+xml',''),('css','css/styles.css','text/css','')]
xfiles=["cover","titlepage","frontispiece","copyright","dedication","epigraph","map"]+[x for p,_,_,chs in PARTS for x in ([f"part{p:02d}"]+[f"ch{ch:02d}" for ch in chs])]+["glossary","authors-note","about"]
for name in xfiles:manifest.append((name,f"xhtml/{name}.xhtml","application/xhtml+xml",""))
for n in range(1,25):manifest.append((f"img{n:02d}",f"images/{img_meta[n]['file']}","image/jpeg","cover-image" if n==1 else ""))
manifest_xml="\n".join(f'<item id="{i}" href="{h}" media-type="{m}"'+(f' properties="{p}"' if p else '')+'/>' for i,h,m,p in manifest)
spine=["cover","titlepage","frontispiece","copyright","dedication","epigraph","nav","map"]
for p,roman,name,chs in PARTS:spine += [f"part{p:02d}"]+[f"ch{ch:02d}" for ch in chs]
spine += ["glossary","authors-note","about"]
spine_xml="\n".join(f'<itemref idref="{x}"/>' for x in spine)
subjects="\n".join(f"<dc:subject>{html.escape(s)}</dc:subject>" for s in SUBJECTS)
opf=f'''<?xml version="1.0" encoding="utf-8"?><package xmlns="http://www.idpf.org/2007/opf" xmlns:dc="http://purl.org/dc/elements/1.1/" version="3.0" unique-identifier="pub-id" prefix="schema: http://schema.org/"><metadata>
<dc:identifier id="pub-id">urn:uuid:{book_uuid}</dc:identifier><dc:title id="title">{TITLE}</dc:title><dc:title id="subtitle">{SUBTITLE}</dc:title><meta property="title-type" refines="#subtitle">subtitle</meta>
<dc:creator id="creator">{AUTHOR}</dc:creator><meta refines="#creator" property="role" scheme="marc:relators">aut</meta><meta refines="#creator" property="file-as">Dahal, Suraj</meta>
<dc:language>en</dc:language><dc:publisher>{PUBLISHER}</dc:publisher><dc:date>{PUBDATE}</dc:date>{subjects}<dc:description>{html.escape(DESCRIPTION)}</dc:description><meta property="dcterms:modified">{modified}</meta>
<meta property="schema:accessMode">textual</meta><meta property="schema:accessMode">visual</meta><meta property="schema:accessModeSufficient">textual</meta>
<meta property="schema:accessibilityFeature">alternativeText</meta><meta property="schema:accessibilityFeature">tableOfContents</meta><meta property="schema:accessibilityFeature">readingOrder</meta><meta property="schema:accessibilityFeature">structuralNavigation</meta>
<meta property="schema:accessibilityHazard">none</meta><meta property="schema:accessibilitySummary">This reflowable edition includes a structured reading order, nested table of contents, semantic headings, and meaningful alternative text and captions for all illustrations. The story is fully readable as text without the images.</meta>
</metadata><manifest>{manifest_xml}</manifest><spine toc="ncx">{spine_xml}</spine></package>'''
(OEBPS/"content.opf").write_text(opf,encoding="utf-8")

(EPUB/"META-INF").mkdir(parents=True,exist_ok=True)
(EPUB/"mimetype").write_text("application/epub+zip",encoding="ascii")
(EPUB/"META-INF"/"container.xml").write_text('<?xml version="1.0" encoding="UTF-8"?><container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/></rootfiles></container>',encoding="utf-8")
epub_path=OUT/OUT_EPUB
with zipfile.ZipFile(epub_path,"w") as z:
    z.write(EPUB/"mimetype","mimetype",compress_type=zipfile.ZIP_STORED)
    for p in sorted(EPUB.rglob("*")):
        if p.is_file() and p.name!="mimetype":z.write(p,p.relative_to(EPUB).as_posix(),compress_type=zipfile.ZIP_DEFLATED,compresslevel=9)

placements=sorted(placements,key=lambda x:x["asset"])
with open(OUT/"image-placement.csv","w",newline="",encoding="utf-8") as f:
    w=csv.DictWriter(f,fieldnames=["asset","source","placement","after"]);w.writeheader();w.writerows(placements)

wd=Document();sec=wd.sections[0]
sec.top_margin=Inches(.75);sec.bottom_margin=Inches(.75);sec.left_margin=Inches(.8);sec.right_margin=Inches(.8)
for sty,size in [("Normal",11),("Title",26),("Heading 1",20),("Heading 2",16)]:
    wd.styles[sty].font.name="Georgia";wd.styles[sty].font.size=Pt(size)
def center_pic(docu,n,width=6.1):
    p=docu.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run().add_picture(str(OPT/f"ty_{n:02d}.jpg"),width=Inches(width))
    c=docu.add_paragraph(CAPTION[n]);c.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in c.runs:r.italic=True;r.font.size=Pt(8)
center_pic(wd,1,6.2);wd.add_page_break()
p=wd.add_paragraph(TITLE,style="Title");p.alignment=WD_ALIGN_PARAGRAPH.CENTER
for t in [SUBTITLE,"A Novel",AUTHOR]:
    p=wd.add_paragraph(t);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
center_pic(wd,2,5.6);wd.add_page_break()
p=wd.add_paragraph("Copyright",style="Heading 1");p.alignment=WD_ALIGN_PARAGRAPH.CENTER
for t in [rights,rights2,disc,"First edition, 2026","Self-published"]:
    q=wd.add_paragraph(t);q.alignment=WD_ALIGN_PARAGRAPH.CENTER
wd.add_page_break();p=wd.add_paragraph("Dedication",style="Heading 1");p.alignment=WD_ALIGN_PARAGRAPH.CENTER
q=wd.add_paragraph(DEDICATION);q.alignment=WD_ALIGN_PARAGRAPH.CENTER
for r in q.runs:r.italic=True
center_pic(wd,5,4.8);wd.add_page_break()
q=wd.add_paragraph("A person who has been absent and unheard of for twelve years\nmay be presumed dead.\n\n— after the Muluki Civil Code of Nepal (paraphrase)");q.alignment=WD_ALIGN_PARAGRAPH.CENTER
for r in q.runs:r.italic=True
wd.add_page_break();wd.add_paragraph("Table of Contents",style="Heading 1")
for pno,roman,name,chs in PARTS:
    wd.add_paragraph(f"Part {roman}: {name}",style="Heading 2")
    for ch in chs:wd.add_paragraph(f"Chapter {ch}. {CHAPTER_TITLES[ch]}")
for t in ["Glossary","Author's Note","About the Author"]:wd.add_paragraph(t)
wd.add_page_break();wd.add_paragraph("Map",style="Heading 1");center_pic(wd,4,6.1)

def add_run_preserved(dst,src):
    runs,texts=transformed_run_texts(src)
    if not runs:dst.add_run(plain_smart(src));return
    for sr,t in zip(runs,texts):
        rr=dst.add_run(t);rr.bold=bool(sr.bold);rr.italic=bool(sr.italic)
def add_manu(docu,plist,ch=None):
    after={}
    if ch:
        for a in chapter_images[ch]:after.setdefault(chapter_anchor[a],[]).append(a)
    for i,p0 in enumerate(plist):
        t=ptext(p0)
        if not t:continue
        if t=="*":
            q=docu.add_paragraph("⁂");q.alignment=WD_ALIGN_PARAGRAPH.CENTER
        else:
            q=docu.add_paragraph();add_run_preserved(q,p0)
            al=effective_alignment(p0)
            if al is not None:q.alignment=al
            if is_doc_para(p0):
                q.paragraph_format.left_indent=Inches(.35)
                for rr in q.runs:rr.font.name="Courier New";rr.font.size=Pt(9)
            else:q.paragraph_format.first_line_indent=Inches(0 if is_noindent(p0) else .25)
        for a in after.get(i,[]):center_pic(docu,a,5.7)

for pno,roman,name,chs in PARTS:
    wd.add_page_break();p=wd.add_paragraph(f"Part {roman}",style="Heading 1");p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=wd.add_paragraph(name);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;center_pic(wd,6+pno,5.8)
    for ch in chs:
        wd.add_page_break();p=wd.add_paragraph(f"Chapter {ch}",style="Heading 2");p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p=wd.add_paragraph(CHAPTER_TITLES[ch]);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;add_manu(wd,chapter_paras[ch],ch)
wd.add_page_break();wd.add_paragraph("Glossary",style="Heading 1")
if gloss_intro:wd.add_paragraph(gloss_intro)
for term,definition in entries:
    p=wd.add_paragraph();r=p.add_run(term);r.bold=True;p.add_run(" — "+definition)
wd.add_page_break();wd.add_paragraph("Author's Note",style="Heading 1");add_manu(wd,note_paras);center_pic(wd,24,5.7)
wd.add_page_break();wd.add_paragraph("About the Author",style="Heading 1");wd.add_paragraph(ABOUT);center_pic(wd,3,5.7)
docx_out=OUT/OUT_DOCX;wd.save(docx_out)

def sha256(p):
    h=hashlib.sha256()
    with open(p,"rb") as f:
        for b in iter(lambda:f.read(1048576),b""):h.update(b)
    return h.hexdigest()
with open(OUT/"CHECKLIST.md","w",encoding="utf-8") as f:
    f.write("# TWELVE YEARS production checklist\n\n")
    f.write("- Metadata: title/subtitle/author/language/publisher/date/UUID/subjects/description/accessibility metadata included.\n")
    f.write("- Official cover: TY_01 embedded and marked as the EPUB cover image.\n")
    f.write("- Images: 24/24 embedded; TY_01 cover + 23 interior/front/back artworks, all with alt text and captions.\n")
    f.write("- Structure: 4 parts and 40 chapters generated as separate XHTML files.\n")
    f.write("- TOC: nested Parts → Chapters in nav.xhtml; landmarks and NCX included.\n")
    f.write("- Glossary: generated as a semantic definition list.\n")
    f.write("- Author's Note: retained at the end as an afterword; About the Author follows.\n")
    f.write("- Smart punctuation: manuscript straight quotation marks converted contextually; three periods normalized to a single ellipsis.\n")
    f.write(f"- EPUB SHA-256: {sha256(epub_path)}\n- Illustrated DOCX SHA-256: {sha256(docx_out)}\n\n")
    f.write("## Image placement\n\n| # | Source | Placement | Anchor |\n|---:|---|---|---|\n")
    for r in placements:f.write(f"| {r['asset']:02d} | {r['source'].replace('|','/')} | {r['placement'].replace('|','/')} | {r['after'].replace('|','/').replace(chr(10),' ')[:160]} |\n")
print("BUILT",epub_path,docx_out)
